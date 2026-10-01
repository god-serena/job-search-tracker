import io
from pypdf import PdfReader
from docx import Document


def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    lower_name = (filename or "").lower()
    text = ""
    if lower_name.endswith(".txt"):
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="replace")
    elif lower_name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for page in reader.pages:
            page_content = page.extract_text()
            if page_content:
                pages_text.append(page_content)
        text = "\n\n".join(pages_text)
    elif lower_name.endswith(".docx"):
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs.append(" | ".join(row_text))
        text = "\n".join(paragraphs)
    else:
        raise ValueError(
            "Unsupported file format. Please upload a .txt, .pdf, or .docx file."
        )

    cleaned = "\n".join(line.rstrip() for line in text.splitlines())
    return cleaned.strip()
